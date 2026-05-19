import asyncio
import csv
import io
import json
from pathlib import Path
from typing import AsyncGenerator, List, Optional

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook

from schemas import CharacterInput, GenerateRequest, CharacterResult, ModelConfig, PromptConfig
from llm import (
    INTRO_PROMPT,
    PERSONA_PROMPT,
    generate_intro,
    generate_persona,
    get_default_model_config,
    get_key_fingerprint,
    get_resolved_request_info,
)

app = FastAPI(title="角色人设批量生成工具")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health")
async def health():
    return {"ok": True, "service": "chatbot_tool"}


async def process_character(
    char: CharacterInput,
    prompt_config: Optional[PromptConfig] = None,
    model_config: Optional[ModelConfig] = None,
) -> CharacterResult:
    persona_system_prompt = (prompt_config.persona_prompt.strip() if prompt_config else "") or None
    intro_system_prompt = (prompt_config.intro_prompt.strip() if prompt_config else "") or None
    base_url = (model_config.base_url.strip() if model_config else "") or None
    model = (model_config.model.strip() if model_config else "") or None
    api_key = (model_config.api_key.strip() if model_config else "") or None
    result = CharacterResult(
        name=char.name,
        gender=char.gender,
        age=char.age,
        role_type=char.role_type,
        personality=char.personality,
        occupation=char.occupation,
        scenario=char.scenario,
        language=char.language,
    )
    try:
        persona = await generate_persona(
            name=char.name,
            gender=char.gender,
            age=char.age,
            role_type=char.role_type,
            personality=char.personality,
            occupation=char.occupation,
            scenario=char.scenario,
            language=char.language,
            system_prompt=persona_system_prompt,
            base_url=base_url,
            model=model,
            api_key=api_key,
        )
        result.persona_prompt = persona
        result.status = "persona_done"

        intro_data = await generate_intro(
            persona,
            char.language,
            system_prompt=intro_system_prompt,
            base_url=base_url,
            model=model,
            api_key=api_key,
        )
        result.introduction = intro_data.get("introduction", "")
        result.prologue = intro_data.get("prologue", "")
        if not result.introduction and not result.prologue:
            result.status = "intro_error"
            result.error = f"简介/开场白为空，模型原始输出：{intro_data.get('raw', '')}"
            return result
        result.status = "done"
    except Exception as e:
        result.status = "error"
        result.error = str(e)
    return result


@app.post("/api/generate")
async def generate(req: GenerateRequest):
    async def event_stream() -> AsyncGenerator[str, None]:
        total = len(req.characters)
        for i, char in enumerate(req.characters):
            result = await process_character(char, req.prompt_config, req.llm_config)
            event = {"index": i, "total": total, "result": result.model_dump()}
            yield json.dumps(event, ensure_ascii=False) + "\n"

    return StreamingResponse(event_stream(), media_type="application/x-ndjson")


@app.get("/api/prompts/defaults")
async def get_default_prompts():
    return {
        "persona_prompt": PERSONA_PROMPT,
        "intro_prompt": INTRO_PROMPT,
    }


@app.get("/api/model/defaults")
async def get_default_model():
    return get_default_model_config()


@app.post("/api/model/test")
async def test_model(config: ModelConfig):
    base_url = config.base_url.strip() or None
    model = config.model.strip() or None
    api_key = config.api_key.strip() or None
    info = get_resolved_request_info(base_url, model)
    key_info = get_key_fingerprint(api_key)
    try:
        text = await generate_persona(
            name="ConnectionTest",
            gender="",
            age="",
            role_type="",
            personality="brief and practical",
            occupation="",
            scenario="",
            language="English",
            system_prompt="Reply with exactly: ok",
            base_url=base_url,
            model=model,
            api_key=api_key,
        )
        return {
            "ok": True,
            **info,
            **key_info,
            "message": text,
        }
    except Exception as e:
        return {
            "ok": False,
            **info,
            **key_info,
            "error": str(e),
        }


COLUMN_MAP = {
    "姓名": "name", "name": "name",
    "性别": "gender", "gender": "gender",
    "年龄": "age", "age": "age",
    "角色类型": "role_type", "role_type": "role_type",
    "核心画像": "personality", "性格画像": "personality", "personality": "personality",
    "职业": "occupation", "occupation": "occupation",
    "初始场景": "scenario", "场景": "scenario", "scenario": "scenario",
    "语言": "language", "language": "language",
}


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    content = await file.read()
    filename = (file.filename or "").lower()
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
    else:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = next(iterator, [])
        rows = []
        for values in iterator:
            rows.append({
                str(header or ""): value
                for header, value in zip(headers, values)
            })

    characters = []
    for r in rows:
        normalized = {
            COLUMN_MAP.get(str(key).strip().lower(), str(key).strip().lower()): value
            for key, value in r.items()
        }
        characters.append({
            "name": "" if normalized.get("name") is None else str(normalized.get("name", "")),
            "gender": "" if normalized.get("gender") is None else str(normalized.get("gender", "")),
            "age": "" if normalized.get("age") is None else str(normalized.get("age", "")),
            "role_type": "" if normalized.get("role_type") is None else str(normalized.get("role_type", "")),
            "personality": "" if normalized.get("personality") is None else str(normalized.get("personality", "")),
            "occupation": "" if normalized.get("occupation") is None else str(normalized.get("occupation", "")),
            "scenario": "" if normalized.get("scenario") is None else str(normalized.get("scenario", "")),
            "language": "" if normalized.get("language") is None else str(normalized.get("language", "English")),
        })
    return {"characters": characters}


@app.post("/api/export")
async def export_results(results: List[dict]):
    headers = [
        "姓名",
        "性别",
        "年龄",
        "角色类型",
        "核心画像",
        "职业",
        "初始场景",
        "语言",
        "人设Prompt",
        "角色简介",
        "开场白",
        "状态",
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "results"
    worksheet.append(headers)
    for r in results:
        worksheet.append([
            r.get("name", ""),
            r.get("gender", ""),
            r.get("age", ""),
            r.get("role_type", ""),
            r.get("personality", ""),
            r.get("occupation", ""),
            r.get("scenario", ""),
            r.get("language", ""),
            r.get("persona_prompt", ""),
            r.get("introduction", ""),
            r.get("prologue", ""),
            r.get("status", ""),
        ])
    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=results.xlsx"},
    )


# Serve frontend static files in production
STATIC_DIR = Path(__file__).resolve().parent / "static"
if STATIC_DIR.exists():
    from fastapi.responses import FileResponse

    @app.get("/{full_path:path}")
    async def serve_frontend(full_path: str):
        file_path = STATIC_DIR / full_path
        if file_path.is_file():
            return FileResponse(file_path)
        return FileResponse(STATIC_DIR / "index.html")
